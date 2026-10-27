from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from app.models.ctRSD_simulator_210 import RSD_sim, FILEPATH_DOMAINS

EXPECTED_HEADERS = ["name", "Rz", "L", "term", "hp5", "prom", "eI", "eO", "s",
                    "invert", "invL", "agL", "AGiloop", "otype", "rna", "us", "ds",
                    "temp_len", "cp", "n", "c", "d", "CDS", "rflap"]

def process_workbook(file : UploadFile):
    # Load the workbook
    workbook = load_workbook(file.file)
    
    # Process sheets
    processed_sheets = [process_sheet(sheet) for sheet in workbook.worksheets]

    # Filter out None values (sheets that didn't match the expected headers)
    data = [sheet for sheet in processed_sheets if sheet is not None]

    # If no valid sheets were found, raise an error
    if not data:
        raise ValueError("No valid sheets found in the workbook. Ensure that the sheets have the expected headers: " + ", ".join(EXPECTED_HEADERS))

    # Make sure all sheets have the same length of data
    lengths = [len(sheet) for sheet in processed_sheets if sheet is not None]
    if len(set(lengths)) > 1:
        raise ValueError("Sheets in the workbook have different lengths of data. Ensure that all sheets have the same number of rows.")

    return data, lengths[0]

def process_sheet(sheet: Workbook):
    # Check if the sheet has the expected headers
    header = [cell.value for cell in sheet[1]]
    if header != EXPECTED_HEADERS:
        return None
    
    # Iterate through the rows in the sheet
    data = [dict(zip(EXPECTED_HEADERS, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]

    return data

def compile_workbook_data(data, num_rows, domains_file=FILEPATH_DOMAINS, get_genbank=0):
    model = RSD_sim()
    model.load_domains(domains_file)

    name = []
    dna_template = []
    rna_template = []
    dna_part_IDs = []
    rna_part_IDs = []
    dna_part_types = []
    rna_part_types = []
    dna_part_orientations = []
    rna_part_orientations = []
    genbank_dna = [] if get_genbank else None
    genbank_rna = [] if get_genbank else None

    unique_parts = {}  # To track unique part IDs across all sheets

    for row_idx in range(num_rows):
        name_row = []
        dna_template_row = []
        rna_template_row = []
        dna_part_IDs_row = []
        rna_part_IDs_row = []
        dna_part_types_row = []
        rna_part_types_row = []
        dna_part_orientations_row = []
        rna_part_orientations_row = []
        genbank_dna_row = []
        genbank_rna_row = []

        for sheet_idx in range(len(data)):
            row = data[sheet_idx][row_idx]

            if row.get("name") is None or row.get("name").lower() == "none":
                continue  # Skip rows without a name

            result = model.ctRSD_seq_compile(
                name=row["name"],
                Rz=row["Rz"] or "Ro",
                L=row["L"] or "L",
                term=row["term"] or "T7t",
                hp5=row["hp5"] or "5hp",
                prom=row["prom"] or "P70a",
                eI=row["eI"] or '',
                eO=row["eO"] or '',
                s=row["s"] or '',
                invert=row["invert"] or 0,
                invL=row["invL"] or 'A',
                agL=row["agL"] or 'TA',
                AGiloop=row["AGiloop"] or 5,
                otype=row["otype"] or 1,
                rna=row["rna"] or 0,
                us=[x.strip()[1:-1] for x in row["us"][1:-1].split(',')] if row["us"] else [],
                ds=[x.strip()[1:-1] for x in row["ds"][1:-1].split(',')] if row["ds"] else [],
                temp_len=row["temp_len"] or 0,
                cp=row["cp"] or '',
                n=row["n"] or '',
                c=row["c"] or 0,
                d=row["d"] or '',
                CDS=row["CDS"] or '',
                rflap=row["rflap"] or '',
                get_genbank=get_genbank
            )

            if result == 'Function terminated, see error.':
                raise ValueError(f"Compilation failed for row {row_idx + 1} in sheet {sheet_idx + 1}. Please check the input parameters.")
            
            name_row.append(row["name"])
            dna_template_row.append(result[0])
            rna_template_row.append(result[1])
            dna_part_IDs_row.append(result[2])
            rna_part_IDs_row.append(result[3])
            dna_part_types_row.append(result[4])
            rna_part_types_row.append(result[5])
            dna_part_orientations_row.append(result[6])
            rna_part_orientations_row.append(result[7])
            genbank_dna_row.append(result[8]) if get_genbank else None
            genbank_rna_row.append(result[9]) if get_genbank else None

            for i in range(len(result[3])):
                part_id = result[3][i]
                part_type = result[5][i]
                part_sequence = model.domain_to_partID_to_seq_d.get(part_type, {}).get(part_id, '')
                unique_parts[part_id] = [part_type, part_sequence]  # Add RNA parts to the set of unique parts

        name.append(name_row)
        dna_template.append(dna_template_row)
        rna_template.append(rna_template_row)
        dna_part_IDs.append(dna_part_IDs_row)
        rna_part_IDs.append(rna_part_IDs_row)
        dna_part_types.append(dna_part_types_row)
        rna_part_types.append(rna_part_types_row)
        dna_part_orientations.append(dna_part_orientations_row)
        rna_part_orientations.append(rna_part_orientations_row)
        genbank_dna.append(genbank_dna_row) if get_genbank else None
        genbank_rna.append(genbank_rna_row) if get_genbank else None
    
    return {
        "name": name,
        "dna_template": dna_template,
        "rna_template": rna_template,
        "dna_part_IDs": dna_part_IDs,
        "rna_part_IDs": rna_part_IDs,
        "dna_part_types": dna_part_types,
        "rna_part_types": rna_part_types,
        "dna_part_orientations": dna_part_orientations,
        "rna_part_orientations": rna_part_orientations,
        "genbank_dna": genbank_dna,
        "genbank_rna": genbank_rna,
        "unique_parts": [{"part_id": k, "part_type": v[0], "part_sequence": v[1]} for k, v in unique_parts.items()]
    }
